"""Post-process pandoc docx: table borders, compact table font, centred images."""
import sys
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.enum.text import WD_ALIGN_PARAGRAPH
d = Document(sys.argv[1])
for t in d.tables:
    tblPr = t._tbl.tblPr
    b = OxmlElement('w:tblBorders')
    for e in ['top', 'bottom', 'insideH']:
        el = OxmlElement(f'w:{e}'); el.set(qn('w:val'), 'single'); el.set(qn('w:sz'), '4'); el.set(qn('w:color'), '808080'); b.append(el)
    tblPr.append(b)
    for w in tblPr.findall(qn('w:tblW')): tblPr.remove(w)
    tw = OxmlElement('w:tblW'); tw.set(qn('w:w'), '5000'); tw.set(qn('w:type'), 'pct'); tblPr.append(tw)
    for row in t.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                for r in p.runs:
                    r.font.size = Pt(8.5)
for p in d.paragraphs:
    if p._p.xpath('.//w:drawing'):
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
from docx.shared import Inches
sec = d.sections[0]
tw = (sec.page_width or Inches(8.5)) - (sec.left_margin or Inches(1)) - (sec.right_margin or Inches(1))
maxh = Inches(8.6)
for shp in d.inline_shapes:
    asp = shp.height / shp.width
    w = min(tw, int(maxh / asp))
    shp.width = int(w); shp.height = int(w * asp)
d.save(sys.argv[1])
